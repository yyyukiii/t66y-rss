#!/usr/bin/env python3
"""Convert generated post-cache JSON files into Excel title lists."""

from __future__ import annotations

import argparse
import glob
import json
import os
import tempfile
import zipfile
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


DEFAULT_PATTERN = "data/**/*_posts_cache.json"
DEFAULT_OUTPUT_SUFFIX = "_titles.xlsx"
HEADER_FILL = "17365D"
ZEBRA_FILL = "DCE6F1"
TITLE_WIDTH = 80


def load_entries(
    json_path: Path,
    title_field: str = "title",
    url_field: str = "url",
) -> list[tuple[str, str | None]]:
    """Read a JSON object/list and return title and URL pairs in source order."""
    with json_path.open("r", encoding="utf-8") as source:
        data = json.load(source)

    if isinstance(data, dict):
        items = data.values()
    elif isinstance(data, list):
        items = data
    else:
        raise ValueError(f"{json_path}: JSON 顶层结构必须是字典或列表")

    entries: list[tuple[str, str | None]] = []
    for item in items:
        if isinstance(item, dict):
            raw_title = item.get(title_field)
            raw_url = item.get(url_field)
        else:
            raw_title = item
            raw_url = None

        url = str(raw_url).strip() if raw_url not in (None, "") else None
        title = str(raw_title).strip() if raw_title not in (None, "") else ""
        entries.append((title or url or "", url))

    return entries


def safe_excel_text(value: str) -> str:
    """Prevent scraped titles from being interpreted as Excel formulas."""
    if value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def safe_hyperlink(url: str | None) -> str | None:
    """Only emit ordinary web links as workbook hyperlinks."""
    if not url:
        return None
    parsed = urlparse(url)
    if parsed.scheme.lower() not in {"http", "https"} or not parsed.netloc:
        return None
    return url


def output_path_for(json_path: Path, suffix: str = DEFAULT_OUTPUT_SUFFIX) -> Path:
    name = json_path.name
    if name.endswith("_posts_cache.json"):
        stem = name[: -len("_posts_cache.json")]
    else:
        stem = json_path.stem
    return json_path.with_name(f"{stem}{suffix}")


def export_to_excel(entries: list[tuple[str, str | None]], output_path: Path) -> None:
    """Write one styled, validated XLSX title list using an atomic replace."""
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "标题清单"

    header = sheet.cell(row=1, column=1, value="标题")
    header.font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
    header.fill = PatternFill("solid", fgColor=HEADER_FILL)
    header.alignment = Alignment(horizontal="center", vertical="center")

    zebra_fill = PatternFill("solid", fgColor=ZEBRA_FILL)
    for row_number, (title, url) in enumerate(entries, start=2):
        cell = sheet.cell(row=row_number, column=1, value=safe_excel_text(title))
        cell.alignment = Alignment(vertical="center")
        hyperlink = safe_hyperlink(url)
        if hyperlink:
            cell.hyperlink = hyperlink
            cell.style = "Hyperlink"
        if row_number % 2 == 0:
            cell.fill = zebra_fill

    sheet.column_dimensions["A"].width = TITLE_WIDTH
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:A{max(1, len(entries) + 1)}"

    fd, temp_name = tempfile.mkstemp(
        prefix=f".{output_path.stem}-",
        suffix=".xlsx",
        dir=output_path.parent,
    )
    os.close(fd)
    temp_path = Path(temp_name)
    try:
        workbook.save(temp_path)
        validate_workbook(temp_path, len(entries))
        os.replace(temp_path, output_path)
    finally:
        workbook.close()
        if temp_path.exists():
            temp_path.unlink()


def validate_workbook(path: Path, entry_count: int) -> None:
    """Check XLSX ZIP integrity and the expected sheet/row count."""
    with zipfile.ZipFile(path) as archive:
        broken_member = archive.testzip()
        if broken_member:
            raise RuntimeError(f"XLSX 压缩包损坏：{broken_member}")

    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if workbook.sheetnames != ["标题清单"]:
            raise RuntimeError(f"工作表异常：{workbook.sheetnames}")
        sheet = workbook["标题清单"]
        if sheet["A1"].value != "标题":
            raise RuntimeError("XLSX 表头验证失败")
        expected_rows = entry_count + 1
        if sheet.max_row != expected_rows:
            raise RuntimeError(
                f"XLSX 行数验证失败：期望 {expected_rows}，实际 {sheet.max_row}"
            )
    finally:
        workbook.close()


def convert_files(pattern: str, suffix: str = DEFAULT_OUTPUT_SUFFIX) -> list[Path]:
    json_paths = sorted(Path(path) for path in glob.glob(pattern, recursive=True))
    if not json_paths:
        raise FileNotFoundError(f"没有找到匹配的 JSON：{pattern}")

    outputs: list[Path] = []
    for json_path in json_paths:
        entries = load_entries(json_path)
        output_path = output_path_for(json_path, suffix)
        export_to_excel(entries, output_path)
        outputs.append(output_path)
        print(f"已生成 {output_path}（{len(entries)} 条）")
    return outputs


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="把 JSON 中的 title 转成 XLSX，并将 url 设为标题超链接。"
    )
    parser.add_argument(
        "pattern",
        nargs="?",
        default=DEFAULT_PATTERN,
        help=f"JSON 路径或 glob（默认：{DEFAULT_PATTERN}）",
    )
    parser.add_argument(
        "--suffix",
        default=DEFAULT_OUTPUT_SUFFIX,
        help=f"输出文件名后缀（默认：{DEFAULT_OUTPUT_SUFFIX}）",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    outputs = convert_files(args.pattern, args.suffix)
    print(f"转换完成，共生成 {len(outputs)} 个 XLSX 文件。")


if __name__ == "__main__":
    main()
