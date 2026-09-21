import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from json_to_xlsx import convert_files, load_entries, output_path_for


class JsonToXlsxTests(unittest.TestCase):
    def test_dictionary_cache_is_converted_in_source_order(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "示例_posts_cache.json"
            source.write_text(
                json.dumps(
                    {
                        "https://example.com/1": {
                            "title": "第一条",
                            "url": "https://example.com/1",
                        },
                        "https://example.com/2": {
                            "title": "=不应成为公式",
                            "url": "javascript:alert(1)",
                        },
                        "https://example.com/3": {
                            "title": "",
                            "url": "https://example.com/3",
                        },
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            outputs = convert_files(str(source))
            self.assertEqual(outputs, [root / "示例_titles.xlsx"])

            workbook = load_workbook(outputs[0], data_only=False)
            try:
                sheet = workbook["标题清单"]
                self.assertEqual(sheet.max_row, 4)
                self.assertEqual(sheet["A2"].value, "第一条")
                self.assertEqual(sheet["A2"].hyperlink.target, "https://example.com/1")
                self.assertEqual(sheet["A3"].value, "'=不应成为公式")
                self.assertIsNone(sheet["A3"].hyperlink)
                self.assertEqual(sheet["A4"].value, "https://example.com/3")
                self.assertEqual(sheet.freeze_panes, "A2")
                self.assertEqual(sheet.auto_filter.ref, "A1:A4")
            finally:
                workbook.close()

    def test_list_json_and_output_naming(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "items.json"
            source.write_text(
                '[{"title":"标题","url":"https://example.com"}]',
                encoding="utf-8",
            )
            self.assertEqual(load_entries(source), [("标题", "https://example.com")])
            self.assertEqual(output_path_for(source), root / "items_titles.xlsx")


if __name__ == "__main__":
    unittest.main()
